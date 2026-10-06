from pathlib import Path
import shutil

from docx import Document
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent
source = ROOT / "Tasks_jul21.docx"
backup = ROOT / "Tasks_jul21_original.docx"
output = ROOT / "Tasks_jul21_updated.docx"
if not backup.exists():
    shutil.copy2(source, backup)

doc = Document(source)
doc.add_page_break()

title = doc.add_paragraph(style="Title")
title.add_run("Response to July Research Tasks")
subtitle = doc.add_paragraph()
run = subtitle.add_run("Updated 11 September 2026")
run.italic = True
run.font.color.rgb = RGBColor(89, 89, 89)

intro = doc.add_paragraph()
intro.add_run("Purpose. ").bold = True
intro.add_run(
    "This section records the response to the July questions and the additional work completed since then. "
    "The current release is a conservative methods and stress test paper. It validates the SURD workflow and tests whether apparent long lag features survive matched controls. It does not claim causal discovery or physical gas transport from these observational data."
)

def heading(text):
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 1"]
    p.add_run(text)
    return p

def status_line(status, color, text):
    p = doc.add_paragraph()
    s = p.add_run(status + ". ")
    s.bold = True
    s.font.color.rgb = RGBColor(*color)
    p.add_run(text)
    return p

def bullet(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(18)
    p.add_run("• ")
    p.add_run(text)
    return p

heading("Exact calculation and normalization")
status_line("COMPLETED", (0, 128, 0), "The adopted pipeline defines the target, predictors, lag convention, histogram construction, SURD atoms, normalized leakage, and sample support explicitly. The implementation is in SURD/utils/surd.py and the release runners record settings and source hashes.")
bullet("Common lag decomposition exports all eleven atoms and normalized leakage for every target in both four variable sets.")
bullet("The four target outputs cover F5100 with the three line components and integrated H beta with the three line components.")
bullet("Analytic unique, redundant, synergy, independence, and additivity checks pass to numerical precision.")

heading("Common lags, unequal lags, and all targets")
status_line("COMPLETED WITH A LIMIT", (0, 102, 204), "Common lag scans and corrected unequal lag scans were executed for every target. Unequal lag results remain descriptive because a matched configuration family null calibration has not been completed.")
bullet("The corrected unequal lag runner uses target relative response phases and permits signed offsets.")
bullet("The adopted velocity intervals are the intervals that generated the release outputs, with boundaries at minus 6000, minus 2000, plus 2000, and plus 6000 kilometres per second.")
bullet("The five unresolved profiles are excluded from the adopted sample and retained in the provenance record.")

heading("Optical continuum and additional wavebands")
status_line("COMPLETED FOR THE AVAILABLE CAMPAIGN", (0, 128, 0), "The optical continuum at 5100 Angstrom is included as the primary driver and as a sensitivity input. The historical NGC 5548 material used here does not provide a sufficiently matched ultraviolet and X ray series for the same adopted pipeline.")
bullet("The absence of ultraviolet and X ray data is now stated as a scope limit rather than filled with unsupported proxies.")
bullet("A larger multi waveband study remains an optional follow up and is not required for this conservative release.")

heading("Significance, power, and intended atoms")
status_line("COMPLETED", (0, 128, 0), "The independent null evaluation was increased to 200 calibration and 200 evaluation realizations. The central power grid contains 3400 realizations.")
bullet("The false positive estimate is 0.075 with a 95 percent Wilson interval from 0.046 to 0.120.")
bullet("Six one factor sensitivity grids vary response width, process timescale, and measurement noise.")
bullet("The synthetic validation uses 100 realizations and identifies the intended atom in 99 to 100 percent of the isolated positive controls. The additive two delay stress test remains difficult and is reported as such.")
bullet("The parameter A is a linear mixing coefficient and is not a variance percentage.")

heading("Autocorrelation, temporal neighbours, and blocked validation")
status_line("COMPLETED WITH BOUNDED INTERPRETATION", (0, 102, 204), "Red noise, irregular sampling, interpolation, and history conditioning are included in the controls. A temporal neighbour exclusion sensitivity was added for one, three, and five day windows.")
bullet("The exclusion study changes the continuous estimator magnitude, which demonstrates estimator sensitivity and does not create evidence for transport.")
bullet("A four block contiguous native date validation was added. Only two blocks satisfy the strict overlap requirement at the fifteen day evaluation lag. Both have negative incremental R squared values.")
bullet("The supported result is nondetection under the declared aggregate null. It neither detects nor refutes physical gas transport.")

heading("What is now ready for release")
status_line("READY FOR CONSERVATIVE RELEASE", (0, 128, 0), "The manuscript, numerical outputs, clean checkout reproduction, unit tests, and PDF visual review are aligned with the adopted pipeline.")
bullet("Sixteen unit tests pass.")
bullet("The clean checkout reproduction matches adopted and sampling outputs byte for byte, with numerical calibration agreement within floating point serialization tolerance.")
bullet("The compiled paper has 22 pages and passed visual review for clipping, overlap, labels, references, headers, footers, and page numbering.")

heading("Optional future extensions")
status_line("NOT RELEASE BLOCKERS", (156, 102, 0), "A larger AGN sample, simultaneous ultraviolet or X ray coverage, and estimator independent atom decomposition would strengthen a future paper but are outside the validated scope of this release.")

for paragraph in doc.paragraphs:
    paragraph.paragraph_format.space_after = Pt(5)

doc.save(output)
print(output)
