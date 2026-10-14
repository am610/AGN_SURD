"""Insert coloured replies into the collaborator's September instructions.

The original document is opened unchanged and each reply is inserted after the
instruction it answers. A status label is green when the requested work is done
and amber when a decision or a review is still needed. Reply text is blue.
"""
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'Instructions for prompt sept26.docx'
OUTPUT = ROOT / 'output' / 'docx' / 'SURD_September_Instructions_With_Replies.docx'
GREEN, AMBER, BLUE = RGBColor(0x17, 0x6B, 0x44), RGBColor(0x94, 0x5B, 0x00), RGBColor(0x17, 0x4A, 0x8B)

# original paragraph index -> (status, label, reply paragraphs)
REPLIES = {
    3: ('amber', 'Process clarification',
        ['The September document is the only specification of the new analysis. The new work is in '
         'collaborator_rebuild (epoch_scan.py, epoch_results, results_section) and does not reuse the interpretation of the '
         'earlier released study, which is unchanged in main.tex. A new chat is a working step; the results themselves cannot '
         'document it.']),
    4: ('green', 'Completed',
        ['All decompositions call the authors’ unchanged SURD routine, pinned to commit 79dbdea85e6754ec2b5457b3e37204c5d53d1815. '
         'The wrapper checks the hashes of both official source files before importing them. The preparation, interpolation, '
         'binning and plotting code is supporting code. No ICCF, JAVELIN or nearest neighbour result enters the analysis. '
         'The authors’ routine was not modified.']),
    5: ('green', 'Completed',
        ['The five series were prepared before the analysis, with python -m collaborator_rebuild.prepare_historical. F5100 comes '
         'from the AGN Watch continuum table (557 observations). Total, blue, core and red H beta are the corrected published '
         'columns of the Wanders and Peterson table (242 spectra on 224 distinct dates). Measurements on the same date are '
         'averaged, with their errors combined in quadrature and divided by the number of measurements. No further scale factor '
         'is applied. F5100 retains its host galaxy contribution. These are observed series, not generated ones.']),
    6: ('green', 'Completed for the historical seasons',
        ['Each series is split into the five historical seasons 1989, 1990, 1991, 1992 and 1993. These are the intervals of '
         'Peterson et al. 2004, Table 2, and they are the Years 1 to 5 of Lu et al. 2022, Table 6. The per season strategy of '
         'Lu and Xi is applied to these data. Their 2015 to 2023 component series are not public, so the analysis uses the '
         'historical seasons, as agreed. The velocity ordering by season shown in this instruction is reproduced in Table 6 of the '
         'results section. Its counterpart for our seasons is the ordering of the peak lags of the unique continuum term for '
         'the blue, core and red targets (Table 5 of the section). It replaces the coarse three bin differences in days, which '
         'must be recalculated with our own velocity definitions.']),
    7: ('green', 'Completed',
        ['Each season has its own five series. Native dates: 1989, 124 continuum and 48 spectra; 1990, 94 and 38; 1991, 65 and 32; '
         '1992, 83 and 46; 1993, 191 and 60. Explicit sampling assumption: the spectra are 3 to 7 days apart in the median, too '
         'sparse for a five dimensional histogram at 30 lags on native dates. Within each season independently, every series is '
         'therefore linearly interpolated onto a daily grid between its observed endpoints, without extrapolation and without '
         'bridging gaps longer than 15 days. This gives 133 to 261 five variable tuples per lag for common lags (118 at the '
         'smallest for the uncommon lags). The tuples are interpolated and autocorrelated, so they are not independent samples.']),
    8: ('green', 'Reference use documented',
        ['The instruction names two PDF files and then three works. We used all three. From the thesis: equations 3.6, 3.16, '
         '3.17 and 3.18 and Figures 4.1 to 4.3. From Xi et al. 2025: Table 3, Section 5.2 and the radius luminosity discussion. '
         'From Lu et al. 2022: Table 6. The references are listed at the end of the results section.']),
    9: ('green', 'Completed',
        ['The deliverable is the LaTeX section collaborator_rebuild/results_section/results_section.tex, not a paper. It contains '
         'eight tables and four example figures; all 120 figures are in three atlas PDFs. The compiled review copy is '
         'output/pdf/surd_epoch_results.pdf (12 pages). It has not been inserted into main.tex.']),
    10: ('amber', 'SURD only; extra diagnostics condensed',
         ['SURD is the only decomposition. Published Lu and Xi results are used for comparison only. The comparison of two and '
          'three histogram bins is a setting of the SURD estimate itself. The categorical benchmark (independent variables give '
          'apparent information and low leakage at small samples) is kept as one table because it affects how leakage is read. '
          'The other diagnostics added in the previous round are no longer part of the main text.']),
    11: ('amber', 'Astrophysical review remains with Ayan',
         ['The astrophysical wording is drafted from Xi et al. 2025 (Table 3, Section 5.2, the radius luminosity discussion) and '
          'Lu et al. 2022 (Table 6). Ayan still has to read both papers and check the wording. This review is not marked '
          'complete. The section does not claim radial inflow or outflow, and does not derive a black hole mass or a radius '
          'from the SURD peaks.']),
    12: ('green', 'Lu et al. 2022 used and cited',
         ['Lu et al. 2022 is cited and used. Table 6 of that paper gives the H beta lag of each of the five historical seasons '
          '(19.7, 18.6, 15.9, 11.0 and 13.0 days). The section compares them with the lag at which the unique continuum '
          'component of the line targets peaks (Table 5 of the section): the SURD lag is longer by a median of 6.0 days for two '
          'bins (5.3 for three), and the rank correlation over the five seasons is 0.67 (0.36 for three bins).']),
    13: ('green', 'Ready for joint reading',
         ['The calculations, tables and figures are ready for reading together. Nothing has been sent to a collaborator or a '
          'journal. The open decisions are listed after the last instruction.']),
    18: ('green', 'Definitions implemented',
         ['Contributions are averaged over target states as in equations 3.16a to 3.16c, ordered as in equation 3.17. Leakage is '
          'conditional target entropy over target entropy (equation 3.6). Normalization follows Figure 4.1: each component over '
          'the joint mutual information between the future target and the whole predictor vector, so the components sum to one. '
          'Leakage has its own denominator and its own panel. Equation 3.18 is applied as described under that instruction.']),
    19: ('green', 'Data location',
         ['The inputs are prepared with python -m collaborator_rebuild.prepare_historical from the AGN Watch tables '
          '(https://www.asc.ohio-state.edu/astronomy/agnwatch/n5548/spectra/ and .../lcv/). The per epoch native series used by '
          'the scans are in collaborator_rebuild/published_results/bins_2/<epoch>/native, which Git ignores. All derived results '
          'are tracked in https://github.com/am610/AGN_SURD under collaborator_rebuild/epoch_results (configuration, hashes, '
          'every component, extrema, atlases).']),
    20: ('green', 'Completed',
         ['Two sets are analysed: total, blue, core, red; and F5100, blue, core, red. Every member is the future target in turn, '
          'and all four members are past predictors, including the history of the target. This is done for each of the five '
          'seasons.']),
    21: ('amber', 'Completed; lag values chosen by us',
         ['The specific mutual information components and the leakage are computed and normalized as in Figure 4.1: every unique, '
          'redundant and synergistic component is divided by the joint mutual information between the future target and the four '
          'predictors, so the 26 components sum to one, and leakage is divided by the target entropy. The two subsets are '
          '{total, blue, core, red} and {F5100, blue, core, red}. The lag placeholders are not filled in the instructions, so these '
          'are our assumptions: base lags of 1 to 30 days in one day steps for both subsets, which covers the 20 to 30 day seasonal '
          'ranges proposed for the modern seasons and the published historical H beta lags of 11 to 20 days. Histograms use two bins '
          '(2^5 = 32 cells, primary) and three bins (3^5 = 243 cells, sensitivity), with quantile boundaries taken from each season\u2019s '
          'own native series. With 133 to 261 tuples the three bin histograms are sparse. Please confirm or replace the lag range.']),
    22: ('green', 'Completed with a counting correction',
         ['The listed combinations give four unique, eleven redundant and eleven synergistic components: 26, not 25. All 26 are '
          'kept. Each figure has U, R and S panels with one curve per component, and a leakage panel, all against lag. There are '
          'no total unique or total synergy curves. The 26 fractions sum to one to within 1.6e-15 in all 7200 decompositions.']),
    23: ('green', 'Completed',
         ['Both sets are run for all four targets in all five seasons: 7200 decompositions, all with defined normalization and '
          'none dropped for support. Table 2 of the paper summarizes the 20 curves (5 seasons, 4 targets) of the total set and Table 3 '
          'the 20 curves of the F5100 set. Figure 3 is an example of the total set; Figures 1, 2 and 4 are examples of the F5100 set. '
          'All 40 common lag curves are in atlas_common.pdf. In the total set the four way redundancy dominates at short lags because '
          'the total is the sum of the components. Mean leakage over lags and targets is 0.61 for the total set and 0.46 for the '
          'F5100 set (two bins).']),
    24: ('amber', 'Completed with a declared interpretation; offsets need confirmation',
         ['Each predictor variable gets its own lag L plus an offset, so the four predictors still give 26 components. Blue '
          'leading samples the blue wing five days further back; red leading does the same for the red wing. Both are run for '
          'both sets, all targets and all seasons. A literal expansion with a second copy of every variable (equation 3.19 '
          'pattern) would have eight predictors and 502 components; it is implemented in lag_design.py but not used. The five '
          'day offset is our choice, of the size of the uncertainties of published H beta lags, and needs confirmation. '
          'Moving one wing back changes mean leakage by at most 0.08 but changes the component allocation by a median total '
          'variation of 0.30.']),
    25: ('green', 'Completed',
         ['There are 120 figures: 5 seasons, 2 sets, 4 targets and 3 lag configurations. Four examples are in the results section; '
          'all 120 are in the three atlas files collaborator_rebuild/epoch_results/figures/atlas_common.pdf, atlas_blue_leading.pdf and '
          'atlas_red_leading.pdf (one page per curve).']),
    26: ('green', 'Completed',
         ['Each season is analysed independently, with bin boundaries taken from its own native series. No season or target '
          'is dropped.']),
    27: ('amber', 'Interpretation supplied; astrophysical review pending',
         ['Leakage is lowest at the shortest lag in 36 of 40 curves (self history) and highest at 20 days or longer in 35 of 40. '
          'At the leakage maximum the largest component is a synergy or the unique continuum term; at the minimum it is the '
          'four way redundancy (total set) or a unique term (F5100 set). In the F5100 set the unique continuum component of the '
          'line targets peaks at 16 to 29 days in 13 of 15 curves, longer than the published 11 to 20 days by a median of 6 '
          'days. The red wing peaks before the blue wing in 1989, 1991, 1992 and 1993 and after it in 1990: red leading blue in '
          'four of five seasons, the pattern Xi et al. report for 2019 and 2021 and attribute to infall. The differences are '
          'five to seven days, one grid step per day and without uncertainty, and the result changes with the bin count in 1990. '
          'Mean leakage across seasons follows the tuple count, as the benchmark predicts, so it is not a physical comparison. '
          'No significance level or flux uncertainty propagation has been computed.']),
}
OPEN_ITEMS = ('amber', 'Open items',
              ['1. Confirm the lag range (1 to 30 days) and the five day offsets for the uncommon lags. '
               '2. Ayan reads Xi et al. 2025 and Lu et al. 2022 and checks the astrophysical wording. '
               '3. Decide whether a significance test and flux uncertainty propagation should be added. '
               '4. Decide whether the literal equation 3.18 expansion with extra copies of each variable is wanted.'])


def style_run(run, color, bold=False):
    run.font.size = Pt(11)
    run.font.color.rgb = color
    run.bold = bold


def new_paragraph_after(anchor, texts):
    """Insert paragraphs after anchor. texts: list of lists of (text, color, bold)."""
    current = anchor
    for runs in texts:
        element = OxmlElement('w:p')
        current._p.addnext(element)
        paragraph = Paragraph(element, current._parent)
        paragraph.paragraph_format.space_after = Pt(6)
        for text, color, bold in runs:
            style_run(paragraph.add_run(text), color, bold)
        current = paragraph
    return current


def main():
    document = Document(str(SOURCE))
    originals = list(document.paragraphs)
    last = {}
    for index, (status, label, paragraphs) in REPLIES.items():
        colour = GREEN if status == 'green' else AMBER
        content = [[('Response ' + label, colour, True)]] + [[(text, BLUE, False)] for text in paragraphs]
        last[index] = new_paragraph_after(originals[index], content)
    status, label, paragraphs = OPEN_ITEMS
    new_paragraph_after(last[27], [[('Response ' + label, AMBER, True)]] + [[(t, BLUE, False)] for t in paragraphs])
    first = originals[0]
    title = OxmlElement('w:p')
    first._p.addprevious(title)
    title_paragraph = Paragraph(title, first._parent)
    title_paragraph.style = document.styles['Title']
    run = title_paragraph.add_run('Detailed replies to the September SURD instructions')
    run.font.color.rgb = RGBColor(0, 0, 0)
    intro = OxmlElement('w:p')
    first._p.addprevious(intro)
    intro_paragraph = Paragraph(intro, first._parent)
    style_run(intro_paragraph.add_run(
        'Responses dated 14 October 2026. The original instructions are retained. Blue text gives our replies; green labels '
        'mark completed work; amber labels mark a decision or a review still needed. All requested calculations have been run. '
        'The lag values and the astrophysical wording need confirmation.'), BLUE)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(OUTPUT))
    answered = len(REPLIES)
    manifest = dict(source='Instructions for prompt sept26.docx',
                    source_sha256=hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
                    output='output/docx/SURD_September_Instructions_With_Replies.docx',
                    output_sha256=hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
                    responses=answered + 1, original_text_and_embedded_material_retained=True,
                    analysis_status='Per season analysis complete; lag values and astrophysical review need confirmation')
    (OUTPUT.parent / 'SURD_September_Replies_manifest.json').write_text(json.dumps(manifest, indent=2) + '\n')
    print(json.dumps(manifest, indent=2))


if __name__ == '__main__':
    main()
